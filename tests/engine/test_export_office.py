"""CVF-01 / CVF-02 / CVF-03: PDF to Word, Excel and PowerPoint."""

from __future__ import annotations

import io
from pathlib import Path

import pymupdf
import pytest
from docx import Document as WordDocument
from openpyxl import load_workbook
from pptx import Presentation
from pptx.util import Emu

from engine.document import Document
from engine.errors import ConversionError, PasswordRequiredError
from engine.export_office import parse_number, to_docx, to_pptx, to_xlsx
from tests import pdfs


@pytest.fixture(scope="module")
def report() -> Document:
    return Document.from_bytes(pdfs.report_pdf())


def _word_text(data: bytes) -> str:
    word = WordDocument(io.BytesIO(data))
    cells = [cell.text for table in word.tables for row in table.rows for cell in row.cells]
    return "\n".join([p.text for p in word.paragraphs] + cells)


@pytest.mark.feature("CVF-01", criterion=1)
def test_a_text_pdf_becomes_a_docx_with_its_text_in_order(report: Document) -> None:
    text = _word_text(to_docx(report))
    assert text.index("Quarterly Report") < text.index("Results") < text.index("Revenue grew") < text.index("Step one")
    assert "Appendix 2" in text and "Closing remarks" in text


@pytest.mark.feature("CVF-01", criterion=2)
def test_a_page_range_limits_the_pages_converted(report: Document) -> None:
    text = _word_text(to_docx(report, [1]))
    assert "Appendix 2" in text and "Quarterly Report" not in text
    with pytest.raises(ConversionError, match="no pages"):
        to_docx(report, [])


@pytest.mark.feature("CVF-01", criterion=3)
def test_pictures_are_carried_into_the_docx(report: Document) -> None:
    word = WordDocument(io.BytesIO(to_docx(report, [0])))
    assert len(word.inline_shapes) >= 1


@pytest.mark.feature("CVF-01", criterion=4)
def test_a_protected_pdf_needs_its_password_to_convert(work_dir: Path) -> None:
    raw = pymupdf.open()
    raw.new_page().insert_text((72, 100), "Secret memo text", fontsize=14)
    path = work_dir / "locked.pdf"
    raw.save(path, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="open-sesame", owner_pw="owner-pw")
    with pytest.raises(PasswordRequiredError):
        Document.open(path)
    unlocked = Document.open(path, password="open-sesame")
    assert "Secret memo text" in _word_text(to_docx(unlocked))


@pytest.mark.feature("CVF-02", criterion=1)
def test_a_ruled_table_becomes_a_sheet_whose_cells_match() -> None:
    result = to_xlsx(Document.from_bytes(pdfs.table_pdf()))
    sheet = load_workbook(io.BytesIO(result.data))["Page 1 Table 1"]
    assert [next(c.value for c in row) for row in sheet.iter_rows()] == [
        "Item",
        "Widget, large",
        'Gadget "pro"',
        "Line break",
    ]
    assert [c.value for c in sheet[1]] == ["Item", "Qty", "Price"]


@pytest.mark.feature("CVF-02", criterion=2)
def test_every_table_gets_its_own_named_sheet_and_the_result_says_where_from() -> None:
    raw = pymupdf.open()
    pdfs.draw_table(raw.new_page(), pdfs.TABLE_ROWS)
    page = raw.new_page()
    pdfs.draw_table(page, [["A", "B"], ["1", "2"]], origin=(72, 100))
    pdfs.draw_table(page, [["C", "D"], ["3", "4"]], origin=(72, 400))
    result = to_xlsx(Document.from_bytes(raw.tobytes()))
    names = load_workbook(io.BytesIO(result.data)).sheetnames
    assert names == ["Page 1 Table 1", "Page 2 Table 1", "Page 2 Table 2"]
    assert [(name, t.page_index + 1, t.number) for name, t in result.sheets] == [
        ("Page 1 Table 1", 1, 1),
        ("Page 2 Table 1", 2, 1),
        ("Page 2 Table 2", 2, 2),
    ]


@pytest.mark.feature("CVF-02", criterion=3)
def test_numbers_are_stored_as_numbers_and_text_stays_text() -> None:
    document = Document.from_bytes(pdfs.table_pdf())
    sheet = load_workbook(io.BytesIO(to_xlsx(document).data))["Page 1 Table 1"]
    assert sheet["B2"].value == 3 and sheet["C2"].value == 1250.0
    assert sheet["C3"].value == -45.5, "an accountant's (45.50) is negative"
    assert sheet["C4"].value == -8
    assert sheet["B4"].value == "007", "a leading zero means an identifier, so it stays text"
    assert sheet["A2"].value == "Widget, large"
    plain = load_workbook(io.BytesIO(to_xlsx(document, numbers=False).data))["Page 1 Table 1"]
    assert plain["B2"].value == "3"


@pytest.mark.parametrize(
    ("text", "value"),
    [
        ("1,250.00", 1250.0),
        ("(45.50)", -45.5),
        ("-8", -8),
        ("12", 12),
        ("0.5", 0.5),
        ("007", None),
        ("$5", None),
        ("12%", None),
        ("1,2", None),
        ("", None),
        ("abc", None),
    ],
)
@pytest.mark.feature("CVF-02", criterion=3)
def test_what_counts_as_a_number(text: str, value: float | None) -> None:
    assert parse_number(text) == value


@pytest.mark.feature("CVF-02", criterion=4)
def test_no_table_is_reported_plainly_not_an_empty_workbook() -> None:
    raw = pymupdf.open()
    raw.new_page().insert_text((72, 100), "A page of prose only.", fontsize=11)
    with pytest.raises(ConversionError, match="no tables were found"):
        to_xlsx(Document.from_bytes(raw.tobytes()))


@pytest.mark.feature("CVF-03", criterion=1)
def test_each_page_becomes_a_slide_matching_its_proportions(report: Document) -> None:
    deck = Presentation(io.BytesIO(to_pptx(report)))
    assert len(deck.slides) == 2
    page = report.raw[0].rect
    assert deck.slide_width / deck.slide_height == pytest.approx(page.width / page.height, rel=0.001)


@pytest.mark.feature("CVF-03", criterion=2)
def test_editable_mode_places_real_text_boxes_at_their_positions(report: Document) -> None:
    deck = Presentation(io.BytesIO(to_pptx(report, [0])))
    slide = deck.slides[0]
    boxes = {shape.text_frame.text.strip(): shape for shape in slide.shapes if shape.has_text_frame}
    title = boxes["Quarterly Report"]
    assert title.text_frame.paragraphs[0].runs[0].font.bold is True
    assert title.text_frame.paragraphs[0].runs[0].font.size.pt == pytest.approx(26, abs=0.5)
    expected_left = Emu(int(72 * 12700 * deck.slide_width / (report.raw[0].rect.width * 12700)))
    assert abs(title.left - expected_left) < Emu(12700 * 8)
    assert abs(title.top - Emu(65 * 12700)) < Emu(12700 * 12)
    assert any(shape.shape_type == 13 for shape in slide.shapes), "the picture came across as a picture"


@pytest.mark.feature("CVF-03", criterion=3)
def test_picture_mode_makes_each_slide_the_rendered_page(report: Document) -> None:
    deck = Presentation(io.BytesIO(to_pptx(report, mode="picture")))
    for slide in deck.slides:
        [shape] = list(slide.shapes)
        assert shape.shape_type == 13
        assert (shape.width, shape.height) == (deck.slide_width, deck.slide_height)


@pytest.mark.feature("CVF-03", criterion=4)
def test_a_page_range_limits_the_slides(report: Document) -> None:
    assert len(Presentation(io.BytesIO(to_pptx(report, [1]))).slides) == 1
    with pytest.raises(ConversionError, match="no pages"):
        to_pptx(report, [])
