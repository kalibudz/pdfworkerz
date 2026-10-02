"""CVT-01..06: Word/Excel/PowerPoint, pictures, HTML, web addresses, text and Markdown to PDF."""

from __future__ import annotations

import functools
import http.server
import io
import threading
import unicodedata
from collections.abc import Iterator
from pathlib import Path

import pymupdf
import pytest
from docx import Document as WordDocument
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation

from engine.errors import ConversionError, MissingToolError
from engine.external import find_tool
from engine.import_docs import (
    html_to_pdf,
    images_to_pdf,
    markdown_to_pdf,
    office_to_pdf,
    text_to_pdf,
    url_to_pdf,
)
from engine.rasters import NamedFile, page_box

needs_libreoffice = pytest.mark.skipif(find_tool("soffice") is None, reason="LibreOffice is not installed")


def _text(pdf: bytes, page: int | None = None) -> str:
    doc = pymupdf.open("pdf", pdf)
    pages = doc if page is None else [doc[page]]
    return "\n".join(p.get_text("text") for p in pages)


def _png(size: tuple[int, int] = (200, 100), colour: tuple[int, int, int] = (20, 140, 60)) -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", size, colour).save(buffer, format="PNG")
    return buffer.getvalue()


# ---------- Word / Excel / PowerPoint ----------


def _docx() -> bytes:
    word = WordDocument()
    word.add_heading("Meeting minutes", 1)
    word.add_paragraph("The committee agreed to proceed with the harbour project.")
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def _xlsx() -> bytes:
    book = Workbook()
    book.active.append(["Region", "Sales"])
    book.active.append(["North", 1250])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _pptx() -> bytes:
    deck = Presentation()
    for title in ("Opening slide", "Second slide", "Closing slide"):
        deck.slides.add_slide(deck.slide_layouts[5]).shapes.title.text = title
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


@needs_libreoffice
@pytest.mark.feature("CVT-01", criterion=1)
def test_a_docx_becomes_a_pdf_with_its_text() -> None:
    text = _text(office_to_pdf(NamedFile("minutes.docx", _docx()), "word"))
    assert "Meeting minutes" in text and "harbour project" in text


@needs_libreoffice
@pytest.mark.feature("CVT-01", criterion=2)
@pytest.mark.parametrize("name", ["minutes.odt", "minutes.rtf"])
def test_other_word_processor_formats_convert_the_same_way(name: str, tmp_path: Path) -> None:
    source = tmp_path / "in.docx"
    source.write_bytes(_docx())
    from engine.external import run_tool

    target = name.split(".")[1]
    run_tool(
        "soffice",
        [
            f"-env:UserInstallation={(tmp_path / 'lo').as_uri()}",
            "--headless",
            "--convert-to",
            target,
            "--outdir",
            tmp_path,
            source,
        ],
        timeout=300,
    )
    produced = tmp_path / f"in.{target}"
    assert produced.exists()
    assert "harbour project" in _text(office_to_pdf(NamedFile(name, produced.read_bytes()), "word"))


@pytest.mark.feature("CVT-01", criterion=3)
@pytest.mark.parametrize(
    ("name", "data"),
    [
        ("notes.docx", b"just some words"),
        ("notes.docx", _xlsx()),
        ("report.pdf", b"%PDF-1.7"),
        ("legacy.doc", b"not an OLE file"),
    ],
)
def test_a_file_that_is_not_a_word_document_is_refused(name: str, data: bytes) -> None:
    with pytest.raises(ConversionError, match=rf"{name} is not a readable Word document"):
        office_to_pdf(NamedFile(name, data), "word")


@pytest.mark.feature("CVT-01", criterion=4)
def test_missing_libreoffice_is_named_with_how_to_install_it(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    monkeypatch.setenv("PDFWORKERZ_SOFFICE", str(tmp_path / "absent"))
    with pytest.raises(MissingToolError, match=r"LibreOffice is not installed.*winget install"):
        office_to_pdf(NamedFile("a.docx", _docx()), "word")


@needs_libreoffice
@pytest.mark.feature("CVT-02", criterion=1)
def test_an_xlsx_becomes_a_pdf_with_its_cells() -> None:
    text = _text(office_to_pdf(NamedFile("sales.xlsx", _xlsx()), "excel"))
    assert "Region" in text and "North" in text and "1250" in text


@needs_libreoffice
@pytest.mark.feature("CVT-02", criterion=2)
def test_csv_and_other_spreadsheet_formats_convert_too() -> None:
    text = _text(office_to_pdf(NamedFile("sales.csv", b"Region,Sales\nSouth,980\n"), "excel"))
    assert "South" in text and "980" in text


@pytest.mark.feature("CVT-02", criterion=3)
def test_a_file_that_is_not_a_spreadsheet_is_refused() -> None:
    with pytest.raises(ConversionError, match=r"sales\.xlsx is not a readable spreadsheet"):
        office_to_pdf(NamedFile("sales.xlsx", _docx()), "excel")


@needs_libreoffice
@pytest.mark.feature("CVT-03", criterion=1)
def test_a_pptx_becomes_a_pdf_with_a_page_per_slide() -> None:
    pdf = office_to_pdf(NamedFile("deck.pptx", _pptx()), "powerpoint")
    assert pymupdf.open("pdf", pdf).page_count == 3
    assert "Second slide" in _text(pdf, 1)


@needs_libreoffice
@pytest.mark.feature("CVT-03", criterion=2)
def test_other_presentation_formats_convert_too(tmp_path: Path) -> None:
    from engine.external import run_tool

    source = tmp_path / "in.pptx"
    source.write_bytes(_pptx())
    run_tool(
        "soffice",
        [
            f"-env:UserInstallation={(tmp_path / 'lo').as_uri()}",
            "--headless",
            "--convert-to",
            "odp",
            "--outdir",
            tmp_path,
            source,
        ],
        timeout=300,
    )
    pdf = office_to_pdf(NamedFile("deck.odp", (tmp_path / "in.odp").read_bytes()), "powerpoint")
    assert pymupdf.open("pdf", pdf).page_count == 3


@pytest.mark.feature("CVT-03", criterion=3)
def test_a_file_that_is_not_a_presentation_is_refused() -> None:
    with pytest.raises(ConversionError, match=r"deck\.pptx is not a readable presentation"):
        office_to_pdf(NamedFile("deck.pptx", b"PK\x03\x04 broken"), "powerpoint")


# ---------- pictures ----------


@pytest.mark.feature("CVT-04", criterion=1)
def test_several_pictures_become_one_page_each_in_order() -> None:
    files = [
        NamedFile("a.png", _png((200, 100), (255, 0, 0))),
        NamedFile("b.png", _png((100, 200), (0, 0, 255))),
        NamedFile("c.png", _png()),
    ]
    doc = pymupdf.open("pdf", images_to_pdf(files))
    assert doc.page_count == 3
    first, second = doc[0].get_pixmap(dpi=20), doc[1].get_pixmap(dpi=20)
    assert first.pixel(first.width // 2, first.height // 2)[:3][0] > 200, "page one is the red picture"
    assert second.pixel(second.width // 2, second.height // 2)[:3][2] > 200, "page two is the blue one"


@pytest.mark.feature("CVT-04", criterion=2)
def test_a_paper_size_and_margin_place_the_picture_centred_and_undistorted() -> None:
    doc = pymupdf.open("pdf", images_to_pdf([NamedFile("a.png", _png((400, 200)))], paper="letter", margin=36))
    page = doc[0]
    assert (page.rect.width, page.rect.height) == (792, 612), "a wide picture turns the sheet landscape"
    shown = pymupdf.Rect(page.get_image_info()[0]["bbox"])
    assert shown.width / shown.height == pytest.approx(2.0, rel=0.01)
    assert shown.x0 >= 36 - 0.01 and shown.x1 <= 792 - 36 + 0.01
    assert (shown.x0 + shown.x1) / 2 == pytest.approx(396, abs=0.5) and (shown.y0 + shown.y1) / 2 == pytest.approx(
        306, abs=0.5
    )
    fit, _ = page_box((100, 100), "fit", 0, dpi=72)
    assert (fit.width, fit.height) == (100, 100)


@pytest.mark.feature("CVT-04", criterion=2)
def test_a_huge_picture_does_not_make_a_huge_page() -> None:
    doc = pymupdf.open("pdf", images_to_pdf([NamedFile("big.png", _png((6000, 4000)))]))
    assert max(doc[0].rect.width, doc[0].rect.height) <= 1190.5


@pytest.mark.feature("CVT-04", criterion=3)
def test_a_sideways_phone_photo_lands_upright() -> None:
    upright = Image.new("RGB", (300, 200), (255, 255, 255))
    sideways = upright.transpose(Image.Transpose.ROTATE_90)
    exif = Image.Exif()
    exif[0x0112] = 8
    buffer = io.BytesIO()
    sideways.save(buffer, format="JPEG", exif=exif)
    page = pymupdf.open("pdf", images_to_pdf([NamedFile("phone.jpg", buffer.getvalue())]))[0]
    assert page.rect.width > page.rect.height


@pytest.mark.feature("CVT-04", criterion=3)
def test_an_upright_jpeg_is_stored_without_being_compressed_again() -> None:
    buffer = io.BytesIO()
    Image.new("RGB", (120, 80), (10, 90, 200)).save(buffer, format="JPEG", quality=61)
    doc = pymupdf.open("pdf", images_to_pdf([NamedFile("p.jpg", buffer.getvalue())]))
    assert doc.extract_image(doc[0].get_images()[0][0])["image"] == buffer.getvalue()


@pytest.mark.feature("CVT-04", criterion=4)
def test_a_file_that_is_not_a_picture_is_refused_by_name_first() -> None:
    with pytest.raises(ConversionError, match=r"notes\.txt is not a picture"):
        images_to_pdf([NamedFile("a.png", _png()), NamedFile("notes.txt", b"words")])
    with pytest.raises(ConversionError, match="no pictures"):
        images_to_pdf([])


@pytest.mark.feature("CVT-04", criterion=1)
@pytest.mark.parametrize("fmt", ["PNG", "JPEG", "BMP", "GIF", "TIFF", "WEBP"])
def test_each_common_picture_format_is_accepted(fmt: str) -> None:
    buffer = io.BytesIO()
    Image.new("RGB", (40, 30), (200, 10, 10)).save(buffer, format=fmt)
    assert pymupdf.open("pdf", images_to_pdf([NamedFile(f"p.{fmt.lower()}", buffer.getvalue())])).page_count == 1


# ---------- HTML, web, text, Markdown ----------

PAGE = (
    "<html><head><title>Notes</title></head><body><h1>Harbour plan</h1><p>Phase <b>one</b> starts in spring.</p>"
    "<ul><li>Dredging</li><li>Piling</li></ul>"
    "<table border='1'><tr><th>Item</th><th>Cost</th></tr><tr><td>Crane</td><td>40,000</td></tr></table></body></html>"
)


@pytest.mark.feature("CVT-05", criterion=1)
def test_html_is_laid_out_and_paginated() -> None:
    pdf = html_to_pdf(PAGE)
    text = _text(pdf)
    for expected in ("Harbour plan", "Phase one starts in spring.", "Dredging", "Crane", "40,000"):
        assert expected in text
    long_page = "<p>" + "A line of ordinary paragraph text that fills the page. " * 700 + "</p>"
    assert pymupdf.open("pdf", html_to_pdf(long_page)).page_count >= 3


@pytest.fixture
def web() -> Iterator[str]:
    class Handler(http.server.SimpleHTTPRequestHandler):
        def do_GET(self) -> None:
            routes = {
                "/page.html": (
                    "text/html",
                    b"<html><body><h1>Served page</h1><img src='/pic.png'><img src='/missing.png'>"
                    b"<p>Hello web</p></body></html>",
                ),
                "/pic.png": ("image/png", _png((160, 90), (255, 0, 0))),
                "/big": ("text/plain", b"x" * (11 * 1024 * 1024)),
            }
            if self.path in routes:
                kind, body = routes[self.path]
                self.send_response(200)
                self.send_header("Content-Type", kind)
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)
            else:
                self.send_error(404)

        def log_message(self, *_args: object) -> None:
            pass

    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), functools.partial(Handler))
    threading.Thread(target=server.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{server.server_port}"
    server.shutdown()


@pytest.mark.feature("CVT-05", criterion=2)
def test_a_web_address_is_fetched_and_converted_with_its_pictures(web: str) -> None:
    result = url_to_pdf(f"{web}/page.html")
    assert "Served page" in _text(result.pdf) and "Hello web" in _text(result.pdf)
    assert len(pymupdf.open("pdf", result.pdf)[0].get_images()) == 1
    assert any("left out" in note for note in result.notes), "the missing picture is reported, not fatal"


@pytest.mark.feature("CVT-05", criterion=2)
@pytest.mark.parametrize(
    "address", ["file:///etc/passwd", "ftp://example.com/a", "javascript:alert(1)", "notaurl", "http://"]
)
def test_any_other_scheme_is_refused(address: str) -> None:
    with pytest.raises(ConversionError, match="only http:// and https://"):
        url_to_pdf(address)


@pytest.mark.feature("CVT-05", criterion=3)
def test_a_page_size_and_margin_can_be_chosen() -> None:
    letter = pymupdf.open("pdf", html_to_pdf(PAGE, paper="letter", margin=100))
    assert (letter[0].rect.width, letter[0].rect.height) == (612, 792)
    left = min(w[0] for w in letter[0].get_text("words"))
    assert left >= 99
    with pytest.raises(ConversionError, match="no room"):
        html_to_pdf(PAGE, margin=400)


@pytest.mark.feature("CVT-05", criterion=4)
def test_fetching_has_a_size_limit_and_a_failed_fetch_says_why(web: str) -> None:
    with pytest.raises(ConversionError, match="larger than 10 MB"):
        url_to_pdf(f"{web}/big")
    with pytest.raises(ConversionError, match=r"could not be fetched.*404"):
        url_to_pdf(f"{web}/nothing-here")
    with pytest.raises(ConversionError, match="could not be fetched"):
        url_to_pdf("http://127.0.0.1:9/")


@pytest.mark.feature("CVT-06", criterion=1)
def test_text_keeps_its_line_breaks_in_monospace_across_pages() -> None:
    lines = [f"row {n:03d}   value {n * 7:>5}" for n in range(220)]
    doc = pymupdf.open("pdf", text_to_pdf("\n".join(lines)))
    assert doc.page_count >= 3
    text = "\n".join(page.get_text("text") for page in doc)
    assert "row 000   value     0" in text and "row 219" in text
    fonts = {font[3] for page in doc for font in page.get_fonts()}
    assert any("Mono" in name or "Courier" in name for name in fonts)
    first = doc[0].get_text("words")
    assert first[0][4] == "row"


@pytest.mark.feature("CVT-06", criterion=2)
def test_markdown_is_styled() -> None:
    source = (
        "# Release notes\n\nSome **bold** and *italic* and `code`.\n\n- first\n- second\n\n"
        "| A | B |\n|---|---|\n| 1 | 2 |\n\n```\nverbatim()\n```\n"
    )
    doc = pymupdf.open("pdf", markdown_to_pdf(source))
    text = unicodedata.normalize("NFKC", doc[0].get_text("text"))  # the typesetter joins f+i into one ligature glyph
    for expected in ("Release notes", "first", "second", "verbatim()"):
        assert expected in text
    sizes = {
        round(s["size"])
        for b in doc[0].get_text("dict")["blocks"]
        for line in b["lines"]
        for s in line["spans"]
        for _ in [0]
        if s["text"].strip() == "Release notes"
    }
    assert max(sizes) > 15, "the heading is bigger than body text"
    flags = {
        s["text"].strip(): s["flags"]
        for b in doc[0].get_text("dict")["blocks"]
        for line in b["lines"]
        for s in line["spans"]
    }
    assert flags["bold"] & 16 and flags["italic"] & 2


@pytest.mark.feature("CVT-06", criterion=2)
def test_raw_html_in_markdown_is_shown_as_text_not_run() -> None:
    doc = pymupdf.open("pdf", markdown_to_pdf("Look: <img src='http://example.com/x.png'> <script>x</script>"))
    assert "<script>" in doc[0].get_text("text") and not doc[0].get_images()


@pytest.mark.feature("CVT-06", criterion=3)
def test_letters_the_base_font_lacks_still_render() -> None:
    text = "café ñ Привет 你好 αβγ"
    out = pymupdf.open("pdf", text_to_pdf(text))[0].get_text("text")
    for part in ("café", "Привет", "你好", "αβγ"):
        assert part in out
