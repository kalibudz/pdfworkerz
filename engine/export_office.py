"""CVF-01 / CVF-02 / CVF-03: a PDF as a Word, Excel or PowerPoint file.

* **Word** is pdf2docx's layout analysis (paragraphs, tables, pictures, columns).
* **Excel** takes the tables engine.tables finds, one sheet each, storing cells that read as
  numbers as numbers so they can be summed.
* **PowerPoint** makes a slide per page: real text boxes and pictures (editable), or the
  rendered page (exact).

All of them export the document as it is *now*, unsaved edits included.
"""

from __future__ import annotations

import io
import re
import tempfile
import warnings
from collections.abc import Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

import pymupdf
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from pdf2docx import Converter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.slide import Slide
from pptx.util import Emu, Pt

from engine.document import Document
from engine.errors import ConversionError
from engine.fonts.classify import split_subset_tag
from engine.tables import Strategy, Table, find_tables

_EMU_PER_POINT = 12700
_MAX_SLIDE_EMU = 51_206_400  # PowerPoint's own limit: 56 inches
_NUMBER = re.compile(r"^[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?$")
_BRACKETED = re.compile(r"^\((.+)\)$")
_SHEET_FORBIDDEN = re.compile(r"[\\/*?:\[\]]")


# ---------- Word ----------


def to_docx(document: Document, page_indices: Sequence[int] | None = None) -> bytes:
    """The pages (all by default) as a .docx. An unreadable result is a :class:`ConversionError`."""
    if document.page_count == 0:
        raise ConversionError("the document has no pages to convert")
    indices = None if page_indices is None else sorted(set(page_indices))
    if indices == []:
        raise ConversionError("there are no pages to convert")
    with tempfile.TemporaryDirectory(prefix="pdfworkerz-docx-") as folder:
        target = Path(folder) / "out.docx"
        converter = Converter(stream=document.to_bytes())
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                converter.convert(str(target), pages=indices, multi_processing=False)
        except Exception as exc:  # pdf2docx raises plain exceptions for anything it can't lay out
            raise ConversionError(f"the document could not be converted to Word ({exc})") from exc
        finally:
            converter.close()
        if not target.exists():
            raise ConversionError("the document could not be converted to Word")
        return target.read_bytes()


# ---------- Excel ----------


def parse_number(text: str) -> int | float | None:
    """A cell's text as a number, or None if it is text. Reads ``1,250.00``, ``-8`` and
    ``(45.50)`` (an accountant's negative). Leading zeros mean an identifier, not a quantity
    (``007``), so those stay text, as does anything with a currency sign or a percent."""
    cell = text.strip()
    bracketed = _BRACKETED.match(cell)
    negative = bool(bracketed)
    if bracketed:
        cell = bracketed.group(1).strip()
    if not _NUMBER.match(cell):
        return None
    digits = cell.lstrip("+-").replace(",", "")
    if len(digits.split(".")[0]) > 1 and digits.startswith("0"):
        return None
    number = float(digits) if "." in digits else int(digits)
    if cell.startswith("-"):
        number = -number
    return -number if negative else number


def _cell(text: str, numbers: bool) -> str | int | float:
    value = parse_number(text) if numbers else None
    return text if value is None else value


def _sheet_name(table: Table, taken: set[str]) -> str:
    base = _SHEET_FORBIDDEN.sub("", f"Page {table.page_index + 1} Table {table.number}")[:31]
    name, suffix = base, 2
    while name in taken:
        name = f"{base[: 31 - len(str(suffix)) - 1]} {suffix}"
        suffix += 1
    taken.add(name)
    return name


@dataclass(frozen=True)
class WorkbookResult:
    data: bytes
    sheets: list[tuple[str, Table]]
    """Each sheet's name with the table (and so the page and table number) it came from."""


def to_xlsx(
    document: Document,
    page_indices: Sequence[int] | None = None,
    *,
    strategy: Strategy = "auto",
    numbers: bool = True,
) -> WorkbookResult:
    """The tables on the pages as a workbook, one sheet per table. ``numbers=False`` keeps every
    cell as text. A document with no table is a :class:`ConversionError` saying so, never an
    empty workbook."""
    tables = find_tables(document, page_indices, strategy=strategy)
    if not tables:
        raise ConversionError("no tables were found on these pages, so there is nothing to put in a spreadsheet")
    workbook = Workbook()
    workbook.remove(workbook.active)
    taken: set[str] = set()
    sheets = []
    for table in tables:
        name = _sheet_name(table, taken)
        sheet = workbook.create_sheet(name)
        for row in table.rows:
            sheet.append([_cell(text, numbers) for text in row])
        for column, width in enumerate(
            (max((len(r[c]) for r in table.rows if c < len(r)), default=8) for c in range(table.column_count)), start=1
        ):
            sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 8), 60)
        sheets.append((name, table))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return WorkbookResult(buffer.getvalue(), sheets)


# ---------- PowerPoint ----------

Fidelity = Literal["editable", "picture"]


def _family(font: str) -> str:
    """``ABCDEF+Arial-BoldMT`` -> ``Arial``: the name PowerPoint can look up."""
    name = split_subset_tag(font)[1]
    return re.split(r"[-,]", name)[0] or "Arial"


def _slide_size(page: pymupdf.Page) -> tuple[int, int]:
    shrink = min(1.0, _MAX_SLIDE_EMU / (max(page.rect.width, page.rect.height) * _EMU_PER_POINT))
    return int(page.rect.width * _EMU_PER_POINT * shrink), int(page.rect.height * _EMU_PER_POINT * shrink)


def to_pptx(
    document: Document,
    page_indices: Sequence[int] | None = None,
    *,
    mode: Fidelity = "editable",
    picture_dpi: int = 150,
) -> bytes:
    """A slide per page. ``editable`` places each line of text as a text box (font, size, colour,
    bold, italic kept) and each picture as a picture; drawn lines and shapes are not carried
    across -- use ``picture`` for a slide that looks exactly like the page."""
    indices = list(range(document.page_count)) if page_indices is None else list(page_indices)
    if not indices:
        raise ConversionError("there are no pages to convert")
    deck = Presentation()
    first = document.raw[indices[0]]
    deck.slide_width, deck.slide_height = (Emu(v) for v in _slide_size(first))
    blank = deck.slide_layouts[6]
    for index in indices:
        page = document.raw[index]
        slide = deck.slides.add_slide(blank)
        scale = deck.slide_width / (first.rect.width * _EMU_PER_POINT)
        if mode == "picture":
            png = page.get_pixmap(dpi=picture_dpi, alpha=False).tobytes("png")
            slide.shapes.add_picture(io.BytesIO(png), 0, 0, deck.slide_width, deck.slide_height)
            continue
        _place_page(slide, page, scale)
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def _place_page(slide: Slide, page: pymupdf.Page, scale: float) -> None:
    to_emu = lambda points: Emu(int(points * _EMU_PER_POINT * scale))  # noqa: E731
    shown = page.rotation_matrix
    data = page.get_text("dict", sort=True, flags=pymupdf.TEXTFLAGS_DICT | pymupdf.TEXT_PRESERVE_IMAGES)
    for block in data["blocks"]:
        box = (pymupdf.Rect(block["bbox"]) * shown).normalize()
        if block["type"] == 1:
            slide.shapes.add_picture(
                io.BytesIO(block["image"]), to_emu(box.x0), to_emu(box.y0), to_emu(box.width), to_emu(box.height)
            )
            continue
        for line in block["lines"]:
            spans = [s for s in line["spans"] if s["text"].strip()]
            if not spans:
                continue
            area = (pymupdf.Rect(line["bbox"]) * shown).normalize()
            frame = slide.shapes.add_textbox(
                to_emu(area.x0), to_emu(area.y0), Emu(to_emu(area.width) + _EMU_PER_POINT * 4), to_emu(area.height)
            )
            text = frame.text_frame
            text.word_wrap = False
            text.margin_left = text.margin_right = text.margin_top = text.margin_bottom = 0
            paragraph = text.paragraphs[0]
            for span in line["spans"]:
                run = paragraph.add_run()
                run.text = span["text"]
                run.font.size = Pt(span["size"] * scale)
                run.font.name = _family(span["font"])
                run.font.bold = bool(span["flags"] & 16)
                run.font.italic = bool(span["flags"] & 2)
                run.font.color.rgb = RGBColor.from_string(f"{span['color']:06X}")
